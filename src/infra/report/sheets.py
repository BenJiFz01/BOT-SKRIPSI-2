"""sheets.py — Fungsi pembuatan tiap sheet laporan Excel."""

from datetime import datetime

from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.infra.report.styles import (
    C_COUNTER, C_HEADER_BLUE, C_HEADER_GOLD, C_LOSS,
    C_PENDING, C_SUBHEADER, C_WIN,
    auto_border, parse_float, set_cell, write_header_row,
)
from src.utils.time_utils import iso_to_wib_excel


def sheet_ringkasan(wb, stats: dict, records: list[dict]) -> None:
    ws = wb.create_sheet("Ringkasan")

    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 2
    ws.column_dimensions["E"].width = 28
    ws.column_dimensions["F"].width = 20
    ws.row_dimensions[1].height = 36
    ws.merge_cells("A1:F1")
    set_cell(ws, 1, 1, "LAPORAN PERFORMA SISTEM DSS TRADING XAU/USD",
             bold=True, size=15, bg=C_HEADER_GOLD, fg="FFFFFF", align="center")
    ws.merge_cells("A2:F2")
    set_cell(ws, 2, 1,
             f"Dibuat: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  |  Skripsi — M. Dzikri Zen",
             italic=True, align="center", bg="FFF8E7")
    ws.merge_cells("A3:F3")

    r = 4

    def section(title: str) -> None:
        nonlocal r
        ws.merge_cells(f"B{r}:C{r}")
        set_cell(ws, r, 2, title, bold=True, size=11, bg=C_HEADER_BLUE, fg="FFFFFF")
        ws.row_dimensions[r].height = 18
        r += 1

    def data_row(label: str, value, formula: str | None = None, pct: bool = False) -> int:
        nonlocal r
        set_cell(ws, r, 2, label)
        cell = ws.cell(row=r, column=3)
        cell.value     = formula if formula else value
        cell.font      = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if pct:
            cell.number_format = "0.0%"
        elif isinstance(value, float):
            cell.number_format = "0.00"
        row_num = r; r += 1
        return row_num

    def blank() -> None:
        nonlocal r; r += 1

    section("A.  TOTAL SINYAL")
    data_row("Total Sinyal Dihasilkan", stats["total_signals"])
    r_decided   = data_row("Sudah Dievaluasi", stats["decided"])
    data_row("Masih Pending",  stats["pending"])
    data_row("Dibatalkan",     stats["cancelled"])
    blank()

    section("B.  DISTRIBUSI OUTCOME")
    r_wtp1 = data_row("WIN TP1",   stats["win_tp1"])
    r_wtp2 = data_row("WIN TP2",   stats["win_tp2"])
    r_wtp3 = data_row("WIN TP3",   stats["win_tp3"])
    r_twin  = data_row("Total WIN", stats["total_wins"], formula=f"=SUM(C{r_wtp1}:C{r_wtp3})")
    data_row("LOSS", stats["losses"])
    blank()

    section("C.  PERFORMA")
    r_wr = data_row("Win Rate", stats["win_rate_pct"] / 100,
                    formula=f"=IFERROR(C{r_twin}/C{r_decided},0)", pct=True)
    ws.cell(row=r_wr, column=3).number_format = "0.0%"
    wr_cell = ws.cell(row=r_wr, column=3)
    wr_cell.fill = PatternFill("solid", fgColor=(C_WIN if stats["win_rate_pct"] >= 50 else C_LOSS))
    data_row("Average RR Realized",        stats["avg_rr"])
    data_row("Avg Durasi Sinyal (menit)",  stats["avg_duration_m"])
    blank()

    by_mode = stats.get("by_mode", {})
    if by_mode:
        section("D.  TREND vs REVERSAL")
        for mode, d in by_mode.items():
            label = "Reversal" if mode == "REVERSAL" else "Trend Following"
            dec   = d["win"] + d["loss"]
            wr    = round(d["win"] / dec * 100, 1) if dec > 0 else 0.0
            data_row(f"{label} — Total",    d["total"])
            data_row(f"{label} — WIN",      d["win"])
            data_row(f"{label} — LOSS",     d["loss"])
            r_mwr = data_row(f"{label} — Win Rate", wr / 100, pct=True)
            ws.cell(row=r_mwr, column=3).number_format = "0.0%"
            ws.cell(row=r_mwr, column=3).value = wr / 100
        blank()

    auto_border(ws, 4, 2, r - 1, 3)

    # Navigasi cepat
    nav_r = 4
    ws.merge_cells(f"E{nav_r}:F{nav_r}")
    set_cell(ws, nav_r, 5, "NAVIGASI CEPAT", bold=True, size=11,
             bg=C_HEADER_BLUE, fg="FFFFFF", align="center")
    nav_r += 1
    for sheet_name, desc in [
        ("Semua Sinyal",     "Histori semua sinyal"),
        ("Per Timeframe",    "Win rate per TF"),
        ("Per Arah",         "BUY vs SELL"),
        ("Trend vs Counter", "Trend Following vs Counter"),
        ("Akurasi Komponen", "Akurasi tiap indikator"),
        ("Equity Curve",     "Grafik kumulatif WIN/LOSS"),
    ]:
        lc = ws.cell(row=nav_r, column=5, value=f"▶  {sheet_name}")
        lc.hyperlink = f"#{sheet_name}!A1"
        lc.font      = Font(color="1F4E79", bold=True, underline="single", size=11)
        lc.alignment = Alignment(vertical="center")
        ws.cell(row=nav_r, column=6, value=desc).font = Font(italic=True, size=10)
        ws.row_dimensions[nav_r].height = 18
        nav_r += 1

    nav_r += 1
    ws.merge_cells(f"E{nav_r}:F{nav_r}")
    set_cell(ws, nav_r, 5, "KETERANGAN WARNA", bold=True, size=11,
             bg=C_HEADER_BLUE, fg="FFFFFF", align="center")
    nav_r += 1
    for label, color in [
        ("WIN (TP1/TP2/TP3)", C_WIN), ("LOSS", C_LOSS),
        ("PENDING", C_PENDING), ("Counter Trend", C_COUNTER),
    ]:
        set_cell(ws, nav_r, 5, f"  {label}", bg=color)
        ws.merge_cells(f"E{nav_r}:F{nav_r}")
        nav_r += 1
    auto_border(ws, 4, 5, nav_r - 1, 6)

    # Data untuk chart (kolom H:I, font kecil/grey agar tidak mencolok)
    pie_start = 4
    set_cell(ws, pie_start, 8, "Kategori", bold=True, bg=C_SUBHEADER)
    set_cell(ws, pie_start, 9, "Jumlah",   bold=True, bg=C_SUBHEADER)
    pie_data_vals = [
        ("WIN TP1", stats["win_tp1"]), ("WIN TP2", stats["win_tp2"]),
        ("WIN TP3", stats["win_tp3"]), ("LOSS", stats["losses"]),
        ("Pending", stats["pending"]),
    ]
    for i, (lbl, val) in enumerate(pie_data_vals):
        set_cell(ws, pie_start + 1 + i, 8, lbl)
        ws.cell(row=pie_start + 1 + i, column=9, value=val)
    for ri in range(pie_start, pie_start + 6):
        for ci in (8, 9):
            ws.cell(row=ri, column=ci).font = Font(size=9, color="AAAAAA")

    pie = PieChart()
    pie.title = "Distribusi Outcome"
    pie.style = 10; pie.width = 14; pie.height = 10
    pie.add_data(Reference(ws, min_col=9, min_row=pie_start+1, max_row=pie_start+5))
    pie.set_categories(Reference(ws, min_col=8, min_row=pie_start+1, max_row=pie_start+5))
    ws.add_chart(pie, "E14")

    by_tf     = stats.get("by_timeframe", {})
    bar_start = pie_start + 8
    set_cell(ws, bar_start, 8, "Timeframe", bold=True, bg=C_SUBHEADER)
    set_cell(ws, bar_start, 9, "Win Rate %", bold=True, bg=C_SUBHEADER)
    for i, (tf, d) in enumerate(sorted(by_tf.items())):
        dec = d["win"] + d["loss"]
        wr  = round(d["win"] / dec * 100, 1) if dec > 0 else 0.0
        ws.cell(row=bar_start+1+i, column=8, value=tf)
        ws.cell(row=bar_start+1+i, column=9, value=wr)
        for ci in (8, 9):
            ws.cell(row=bar_start+1+i, column=ci).font = Font(size=9, color="AAAAAA")

    if by_tf:
        bar = BarChart()
        bar.type = "col"; bar.grouping = "clustered"; bar.overlap = 0
        bar.title = "Win Rate per Timeframe (%)"
        bar.y_axis.title = "Win Rate (%)"; bar.x_axis.title = "Timeframe"
        bar.style = 10; bar.width = 14; bar.height = 10
        n = len(by_tf)
        bar.add_data(Reference(ws, min_col=9, min_row=bar_start, max_row=bar_start+n), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=8, min_row=bar_start+1, max_row=bar_start+n))
        ws.add_chart(bar, "E28")


def sheet_semua_sinyal(wb, records: list[dict]) -> None:
    ws = wb.create_sheet("Semua Sinyal")
    headers = [
        "No", "Signal ID", "Tanggal (WIB)", "Simbol", "TF", "Arah", "Mode",
        "Entry", "Stop Loss", "TP1", "TP2", "TP3", "RR",
        "Trigger", "Confluence", "HTF Bias",
        "Pattern", "Fibonacci", "SnR", "SnD", "Divergence", "Sesi",
        "Outcome",
        "TP1 Hit", "TP2 Hit", "TP3 Hit", "SL Hit",
        "Durasi (mnt)", "Harga Resolve", "Catatan",
    ]
    write_header_row(ws, 1, headers, bg=C_HEADER_BLUE)
    for i, w in enumerate([4,38,20,9,5,6,14,10,10,10,10,10,6,10,10,22,18,20,16,20,14,10,10,20,20,20,20,12,14,20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    for idx, rec in enumerate(sorted(records, key=lambda x: x.get("timestamp", "")), 1):
        outcome = rec.get("outcome", "PENDING")
        if outcome in ("WIN_TP1", "WIN_TP2", "WIN_TP3"): row_bg = C_WIN
        elif outcome == "LOSS":                            row_bg = C_LOSS
        elif outcome == "PENDING":                         row_bg = C_PENDING
        else:                                              row_bg = C_COUNTER if rec.get("signal_mode") == "REVERSAL" else None

        mode = "Reversal" if rec.get("signal_mode") == "REVERSAL" else "Trend"
        vals = [
            idx, rec.get("signal_id",""), iso_to_wib_excel(rec.get("timestamp","")),
            rec.get("symbol",""), rec.get("timeframe",""), rec.get("direction",""), mode,
            parse_float(rec.get("entry")), parse_float(rec.get("sl")),
            parse_float(rec.get("tp1")),  parse_float(rec.get("tp2")),
            parse_float(rec.get("tp3")),  parse_float(rec.get("rr")),
            rec.get("trigger_score",""), rec.get("confluence_score",""), rec.get("htf_bias",""),
            rec.get("pattern_names",""), rec.get("fib_detail",""),
            rec.get("snr_detail",""),    rec.get("snd_detail",""),
            rec.get("divergence_detail",""), rec.get("session_name",""), outcome,
            iso_to_wib_excel(rec.get("tp1_hit_time","")), iso_to_wib_excel(rec.get("tp2_hit_time","")),
            iso_to_wib_excel(rec.get("tp3_hit_time","")), iso_to_wib_excel(rec.get("sl_hit_time","")),
            int(rec.get("duration_minutes", 0) or 0),
            parse_float(rec.get("outcome_price")), rec.get("notes",""),
        ]
        row_num = idx + 1
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            if row_bg:
                cell.fill = PatternFill("solid", fgColor=row_bg)
            cell.alignment = Alignment(wrap_text=False, vertical="center")

        outcome_col = headers.index("Outcome") + 1
        set_cell(ws, row_num, outcome_col, outcome, bold=True,
                 bg=C_WIN if "WIN" in outcome else (C_LOSS if outcome == "LOSS" else C_PENDING))

    auto_border(ws, 1, 1, len(records) + 1, len(headers))

    leg = len(records) + 3
    ws.merge_cells(f"A{leg}:D{leg}")
    set_cell(ws, leg, 1, "Keterangan Warna:", bold=True)
    for label, color in [("WIN", C_WIN), ("LOSS", C_LOSS), ("PENDING", C_PENDING)]:
        leg += 1
        set_cell(ws, leg, 1, f"  {label}", bg=color)


def sheet_per_tf(wb, stats: dict) -> None:
    ws    = wb.create_sheet("Per Timeframe")
    by_tf = stats.get("by_timeframe", {})
    set_cell(ws, 1, 1, "PERFORMA PER TIMEFRAME", bold=True, size=12, bg=C_HEADER_GOLD, fg="FFFFFF")
    ws.merge_cells("A1:G1")
    headers = ["Timeframe", "Total", "WIN", "LOSS", "Pending", "Win Rate (%)", "Keterangan"]
    write_header_row(ws, 2, headers)
    for i, w in enumerate([12,8,8,8,8,14,20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 3
    for tf, d in sorted(by_tf.items()):
        dec  = d["win"] + d["loss"]
        wr   = round(d["win"] / dec * 100, 1) if dec > 0 else 0.0
        note = "Data cukup" if dec >= 10 else ("Data terbatas" if dec > 0 else "Belum ada data")
        bg   = C_WIN if wr >= 60 else (C_LOSS if dec > 0 and wr < 40 else None)
        for col, val in enumerate([tf, d["total"], d["win"], d["loss"], d["pending"], wr, note], 1):
            set_cell(ws, r, col, val, bg=bg)
        r += 1
    auto_border(ws, 1, 1, r - 1, 7)
    set_cell(ws, r + 1, 1, "* Hijau >= 60% | Merah < 40%", italic=True)


def sheet_per_arah(wb, stats: dict) -> None:
    ws     = wb.create_sheet("Per Arah")
    by_dir = stats.get("by_direction", {})
    set_cell(ws, 1, 1, "PERFORMA PER ARAH (BUY vs SELL)", bold=True, size=12, bg=C_HEADER_GOLD, fg="FFFFFF")
    ws.merge_cells("A1:F1")
    write_header_row(ws, 2, ["Arah", "Total", "WIN", "LOSS", "Pending", "Win Rate (%)"])
    for i, w in enumerate([10,8,8,8,8,14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 3
    for direc, d in by_dir.items():
        dec = d["win"] + d["loss"]
        wr  = round(d["win"] / dec * 100, 1) if dec > 0 else 0.0
        bg  = C_WIN if wr >= 60 else (C_LOSS if dec > 0 and wr < 40 else None)
        for col, val in enumerate([direc, d["total"], d["win"], d["loss"], d["pending"], wr], 1):
            set_cell(ws, r, col, val, bg=bg)
        r += 1
    auto_border(ws, 1, 1, r - 1, 6)


def sheet_trend_vs_counter(wb, stats: dict) -> None:
    ws      = wb.create_sheet("Trend vs Counter")
    by_mode = stats.get("by_mode", {})
    set_cell(ws, 1, 1, "TREND FOLLOWING vs COUNTER TREND", bold=True, size=12, bg=C_HEADER_GOLD, fg="FFFFFF")
    ws.merge_cells("A1:G1")
    write_header_row(ws, 2, ["Mode", "Total", "WIN", "LOSS", "Pending", "Win Rate (%)", "Keterangan"])
    for i, w in enumerate([18,8,8,8,8,14,40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 3
    labels = {"CONTINUATION": "Trend Following", "REVERSAL": "Reversal", "BREAKOUT": "Breakout"}
    notes  = {"CONTINUATION": "Mengikuti arah HTF bias", "REVERSAL": "Melawan HTF bias", "BREAKOUT": "Breakout momentum"}
    for mode, d in by_mode.items():
        dec  = d["win"] + d["loss"]
        wr   = round(d["win"] / dec * 100, 1) if dec > 0 else 0.0
        bg   = C_COUNTER if mode == "REVERSAL" else None
        for col, val in enumerate([labels.get(mode,mode), d["total"], d["win"], d["loss"], d["pending"], wr, notes.get(mode,"")], 1):
            set_cell(ws, r, col, val, bg=bg)
        r += 1
    auto_border(ws, 1, 1, r - 1, 7)


def sheet_akurasi_komponen(wb, stats: dict) -> None:
    ws   = wb.create_sheet("Akurasi Komponen")
    comp = stats.get("component_accuracy", {})
    set_cell(ws, 1, 1, "AKURASI PER KOMPONEN ANALISIS TEKNIKAL", bold=True, size=12, bg=C_HEADER_GOLD, fg="FFFFFF")
    ws.merge_cells("A1:F1")
    set_cell(ws, 2, 1, "Akurasi = % sinyal WIN dari semua sinyal yang memakai komponen ini", italic=True)
    ws.merge_cells("A2:F2")
    write_header_row(ws, 3, ["Komponen", "Layer", "Total", "WIN", "LOSS", "Akurasi (%)"])
    for i, w in enumerate([18,12,12,8,8,14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    layer_map = {
        "EMA200":"Trigger","EMA50":"Trigger","RSI":"Trigger","MACD":"Trigger",
        "PATTERN":"Confluence","DIVERGENCE":"Confluence","FIBONACCI":"Confluence",
        "SNR":"Confluence","SND":"Confluence",
    }
    r = 4
    if comp:
        for name, v in sorted(comp.items(), key=lambda x: -x[1]["accuracy"]):
            bg = C_WIN if v["accuracy"] >= 60 else (C_LOSS if v["total"] > 0 and v["accuracy"] < 40 else None)
            for col, val in enumerate([name, layer_map.get(name,"-"), v["total"], v["win"], v["loss"], v["accuracy"]], 1):
                set_cell(ws, r, col, val, bg=bg)
            r += 1
    else:
        ws.merge_cells(f"A{r}:F{r}")
        set_cell(ws, r, 1, "Belum ada data", italic=True); r += 1
    auto_border(ws, 1, 1, r - 1, 6)
    set_cell(ws, r + 1, 1, "* Hijau >= 60% | Merah < 40%", italic=True)


def sheet_equity_curve(wb, records: list[dict]) -> None:
    ws = wb.create_sheet("Equity Curve")
    set_cell(ws, 1, 1, "EQUITY CURVE — KUMULATIF WIN/LOSS", bold=True, size=12, bg=C_HEADER_GOLD, fg="FFFFFF")
    ws.merge_cells("A1:F1")
    set_cell(ws, 2, 1, "WIN = +1 | LOSS = -1 | Equity = akumulasi sejak sinyal pertama", italic=True)
    ws.merge_cells("A2:F2")
    write_header_row(ws, 3, ["No", "Tanggal Resolve", "Signal ID", "Arah", "Outcome", "Equity"])
    for i, w in enumerate([5,22,38,6,12,10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    decided = sorted(
        [r for r in records if r.get("outcome") not in ("PENDING","CANCELLED","") and r.get("outcome_time")],
        key=lambda x: x.get("outcome_time", ""),
    )
    equity   = 0
    data_row = 4
    for i, rec in enumerate(decided, 1):
        outcome = rec.get("outcome", "")
        if outcome in ("WIN_TP1","WIN_TP2","WIN_TP3"):
            equity += 1; bg = C_WIN
        elif outcome == "LOSS":
            equity -= 1; bg = C_LOSS
        else:
            continue
        for col, val in enumerate([i, iso_to_wib_excel(rec.get("outcome_time","")),
                                    rec.get("signal_id",""), rec.get("direction",""), outcome, equity], 1):
            set_cell(ws, data_row, col, val, bg=bg)
        ws.cell(row=data_row, column=6).font = Font(bold=True)
        data_row += 1

    if data_row == 4:
        ws.merge_cells("A4:F4")
        set_cell(ws, 4, 1, "Belum ada sinyal yang resolve (WIN/LOSS)", italic=True)
        data_row = 5

    auto_border(ws, 1, 1, data_row - 1, 6)
    set_cell(ws, data_row + 1, 1, f"Equity Akhir  : {equity:+d}", bold=True)
    set_cell(ws, data_row + 2, 1, f"Total Resolve : {len(decided)} sinyal", bold=True)

    n_rows = data_row - 4 - 2
    if n_rows > 0:
        chart = LineChart()
        chart.title = "Equity Curve"
        chart.y_axis.title = "Equity"; chart.x_axis.title = "Urutan Sinyal"
        chart.style = 10; chart.width = 22; chart.height = 14; chart.grouping = "standard"
        chart.add_data(Reference(ws, min_col=6, min_row=4, max_row=3+n_rows))
        s = chart.series[0]
        s.title = SeriesLabel(v="Equity")
        s.graphicalProperties.line.solidFill = "1F4E79"
        s.graphicalProperties.line.width     = 25000
        ws.add_chart(chart, "H3")
