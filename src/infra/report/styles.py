"""styles.py — Konstanta warna dan helper styling openpyxl."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

C_HEADER_BLUE = "1F4E79"
C_HEADER_GOLD = "C9A000"
C_WIN         = "E2EFDA"
C_LOSS        = "FCE4D6"
C_PENDING     = "FFF2CC"
C_COUNTER     = "DAE3F3"
C_SUBHEADER   = "D6E4F0"


def set_cell(
    ws, row: int, col: int, value=None,
    bold: bool = False, italic: bool = False, size: int = 11,
    bg: str | None = None, fg: str = "000000",
    align: str = "left",
) -> None:
    """Tulis nilai ke sel dan terapkan styling."""
    cell = ws.cell(row=row, column=col, value=value)
    cell.font      = Font(bold=bold, italic=italic, size=size, color=fg)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=False)
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)


def write_header_row(
    ws, row: int, headers: list[str],
    bg: str = C_HEADER_BLUE, fg: str = "FFFFFF",
) -> None:
    """Tulis baris header dengan styling bold + warna."""
    for col, h in enumerate(headers, 1):
        set_cell(ws, row, col, h, bold=True, bg=bg, fg=fg, align="center")
    ws.row_dimensions[row].height = 18


def auto_border(ws, r1: int, c1: int, r2: int, c2: int) -> None:
    """Terapkan border tipis ke semua sel dalam range."""
    thin   = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
        for cell in row:
            cell.border = border


def parse_float(val) -> float:
    """Parse float dengan aman, return 0.0 jika None/invalid."""
    try:
        return float(val or 0)
    except Exception:
        return 0.0
